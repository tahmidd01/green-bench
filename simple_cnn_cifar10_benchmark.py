# ==============================================================================
# GREEN AI BENCHMARK
# PyTorch vs TensorFlow on CIFAR-10
# ==============================================================================

import os

# Reduce unnecessary TensorFlow console messages
os.environ["TF_CPP_MIN_LOG_LEVEL"] = "2"

import time
import numpy as np
import torch
import torch.nn as nn
import torch.optim as optim
import torchvision
import torchvision.transforms as transforms
import tensorflow as tf

from tensorflow.keras import layers, Model
from codecarbon import OfflineEmissionsTracker
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH


# ==============================================================================
# 0. CONFIGURATION
# ==============================================================================

SEED = 42
BATCH_SIZE = 64
NUM_CLASSES = 10
LEARNING_RATE = 0.001

# Set random seeds for reproducibility
torch.manual_seed(SEED)
np.random.seed(SEED)
tf.random.set_seed(SEED)

print("=" * 70)
print("GREEN AI BENCHMARK")
print("PyTorch vs TensorFlow")
print("=" * 70)

print(f"PyTorch version   : {torch.__version__}")
print(f"TensorFlow version: {tf.__version__}")

# ------------------------------------------------------------------------------
# IMPORTANT:
# Both frameworks are intentionally benchmarked on CPU.
#
# This makes the comparison fair on native Windows, because otherwise PyTorch
# could potentially use CUDA while TensorFlow uses the CPU.
# ------------------------------------------------------------------------------

device = torch.device("cpu")

print(f"Benchmark device  : {device}")
print()


# ==============================================================================
# 1. MODEL DEFINITIONS
# ==============================================================================

class SimpleCNN_PyTorch(nn.Module):

    def __init__(self, num_classes=10):
        super(SimpleCNN_PyTorch, self).__init__()

        self.features = nn.Sequential(

            nn.Conv2d(
                in_channels=3,
                out_channels=32,
                kernel_size=3,
                padding=1
            ),

            nn.BatchNorm2d(32),

            nn.ReLU(),

            nn.MaxPool2d(
                kernel_size=2,
                stride=2
            ),

            nn.Conv2d(
                in_channels=32,
                out_channels=64,
                kernel_size=3,
                padding=1
            ),

            nn.BatchNorm2d(64),

            nn.ReLU(),

            nn.MaxPool2d(
                kernel_size=2,
                stride=2
            )
        )

        self.classifier = nn.Sequential(

            nn.AdaptiveAvgPool2d((1, 1)),

            nn.Flatten(),

            nn.Linear(
                64,
                128
            ),

            nn.ReLU(),

            nn.Linear(
                128,
                num_classes
            )
        )

    def forward(self, x):

        x = self.features(x)

        x = self.classifier(x)

        return x


def SimpleCNN_TensorFlow(num_classes=10):

    inputs = layers.Input(
        shape=(32, 32, 3)
    )

    x = layers.Conv2D(
        filters=32,
        kernel_size=3,
        padding="same"
    )(inputs)

    x = layers.BatchNormalization()(x)

    x = layers.ReLU()(x)

    x = layers.MaxPooling2D(
        pool_size=2,
        strides=2
    )(x)

    x = layers.Conv2D(
        filters=64,
        kernel_size=3,
        padding="same"
    )(x)

    x = layers.BatchNormalization()(x)

    x = layers.ReLU()(x)

    x = layers.MaxPooling2D(
        pool_size=2,
        strides=2
    )(x)

    x = layers.GlobalAveragePooling2D()(x)

    x = layers.Dense(
        128,
        activation="relu"
    )(x)

    outputs = layers.Dense(
        num_classes
    )(x)

    return Model(
        inputs=inputs,
        outputs=outputs,
        name="SimpleCNN_TF"
    )


# ==============================================================================
# 2. VERIFY MODEL PARAMETER PARITY
# ==============================================================================

print("=" * 70)
print("MODEL PARAMETER CHECK")
print("=" * 70)

pt_model_temp = SimpleCNN_PyTorch(
    num_classes=NUM_CLASSES
)

tf_model_temp = SimpleCNN_TensorFlow(
    num_classes=NUM_CLASSES
)

# Count only TRAINABLE PyTorch parameters
pt_params = sum(
    p.numel()
    for p in pt_model_temp.parameters()
    if p.requires_grad
)

# Count only TRAINABLE TensorFlow parameters
#
# We do NOT use tf_model_temp.count_params() here because that also counts
# BatchNormalization moving statistics, which are non-trainable.
tf_params = sum(
    int(np.prod(variable.shape))
    for variable in tf_model_temp.trainable_variables
)

print(f"PyTorch trainable parameters   : {pt_params:,}")
print(f"TensorFlow trainable parameters: {tf_params:,}")

if pt_params == tf_params:

    parity_status = "Match"

    print("Parameter parity: MATCH")

else:

    parity_status = "Mismatch"

    print("Parameter parity: MISMATCH")


# ==============================================================================
# 3. DATASET PREPARATION
# ==============================================================================

print("\n" + "=" * 70)
print("LOADING CIFAR-10 DATASET")
print("=" * 70)


# ------------------------------------------------------------------------------
# PyTorch CIFAR-10
# ------------------------------------------------------------------------------

transform_pt = transforms.Compose([

    transforms.ToTensor(),

    transforms.Normalize(
        mean=(0.4914, 0.4822, 0.4465),
        std=(0.2023, 0.1994, 0.2010)
    )
])

pt_trainset = torchvision.datasets.CIFAR10(

    root="./data",

    train=True,

    download=True,

    transform=transform_pt
)

# Generator makes PyTorch shuffling reproducible
pt_generator = torch.Generator()

pt_generator.manual_seed(SEED)

pt_loader = torch.utils.data.DataLoader(

    pt_trainset,

    batch_size=BATCH_SIZE,

    shuffle=True,

    num_workers=0,

    generator=pt_generator
)

print(f"PyTorch training samples: {len(pt_trainset):,}")


# ------------------------------------------------------------------------------
# TensorFlow CIFAR-10
# ------------------------------------------------------------------------------

(tf_x_train, tf_y_train), _ = tf.keras.datasets.cifar10.load_data()

# Convert labels from shape:
# (50000, 1)
#
# to:
# (50000,)
tf_y_train = tf_y_train.squeeze()

# Use explicit float32 constants
tf_mean = tf.constant(
    [0.4914, 0.4822, 0.4465],
    dtype=tf.float32
)

tf_std = tf.constant(
    [0.2023, 0.1994, 0.2010],
    dtype=tf.float32
)


def preprocess_tf(image, label):

    image = tf.cast(
        image,
        tf.float32
    )

    image = image / 255.0

    image = (
        image - tf_mean
    ) / tf_std

    return image, label


tf_dataset = tf.data.Dataset.from_tensor_slices(
    (
        tf_x_train,
        tf_y_train
    )
)

tf_dataset = tf_dataset.shuffle(
    buffer_size=len(tf_x_train),
    seed=SEED,
    reshuffle_each_iteration=False
)

tf_dataset = tf_dataset.map(
    preprocess_tf,
    num_parallel_calls=tf.data.AUTOTUNE
)

tf_dataset = tf_dataset.batch(
    BATCH_SIZE
)

tf_dataset = tf_dataset.prefetch(
    tf.data.AUTOTUNE
)

print(f"TensorFlow training samples: {len(tf_x_train):,}")

print("Dataset preparation complete.")


# ==============================================================================
# 4. PYTORCH BENCHMARK
# ==============================================================================

print("\n" + "=" * 70)
print("RUNNING PYTORCH BENCHMARK - 1 EPOCH")
print("=" * 70)

pt_model = SimpleCNN_PyTorch(
    num_classes=NUM_CLASSES
).to(device)

criterion = nn.CrossEntropyLoss()

optimizer = optim.Adam(
    pt_model.parameters(),
    lr=LEARNING_RATE
)

tracker_pt = OfflineEmissionsTracker(
    country_iso_code="BGD",
    project_name="PyTorch_SimpleCNN"
)

tracker_pt.start()

start_pt = time.perf_counter()

pt_model.train()

pt_total_loss = 0.0
pt_batches = 0

for batch_index, (inputs, targets) in enumerate(pt_loader):

    inputs = inputs.to(device)

    targets = targets.to(device)

    optimizer.zero_grad()

    outputs = pt_model(inputs)

    loss = criterion(
        outputs,
        targets
    )

    loss.backward()

    optimizer.step()

    pt_total_loss += loss.item()

    pt_batches += 1

    # Show progress every 100 batches
    if (batch_index + 1) % 100 == 0:

        print(
            f"PyTorch batch "
            f"{batch_index + 1}/{len(pt_loader)} "
            f"| Loss: {loss.item():.4f}"
        )

end_pt = time.perf_counter()

emissions_pt = tracker_pt.stop()

pt_duration = end_pt - start_pt

# Safety check in case CodeCarbon returns None
if emissions_pt is None:
    emissions_pt = 0.0

pt_average_loss = (
    pt_total_loss / pt_batches
)

print("\nPyTorch completed.")

print(
    f"Execution time : "
    f"{pt_duration:.4f} seconds"
)

print(
    f"Average loss   : "
    f"{pt_average_loss:.4f}"
)

print(
    f"CO2 emissions  : "
    f"{emissions_pt:.8f} kg CO2eq"
)


# ==============================================================================
# 5. TENSORFLOW BENCHMARK
# ==============================================================================

print("\n" + "=" * 70)
print("RUNNING TENSORFLOW BENCHMARK - 1 EPOCH")
print("=" * 70)

tf_model = SimpleCNN_TensorFlow(
    num_classes=NUM_CLASSES
)

tf_loss_fn = tf.keras.losses.SparseCategoricalCrossentropy(
    from_logits=True
)

tf_optimizer = tf.keras.optimizers.Adam(
    learning_rate=LEARNING_RATE
)

tracker_tf = OfflineEmissionsTracker(
    country_iso_code="BGD",
    project_name="TensorFlow_SimpleCNN"
)

tracker_tf.start()

start_tf = time.perf_counter()

tf_total_loss = 0.0
tf_batches = 0

for batch_index, (x_batch, y_batch) in enumerate(tf_dataset):

    with tf.GradientTape() as tape:

        logits = tf_model(
            x_batch,
            training=True
        )

        loss_value = tf_loss_fn(
            y_batch,
            logits
        )

    grads = tape.gradient(
        loss_value,
        tf_model.trainable_weights
    )

    tf_optimizer.apply_gradients(
        zip(
            grads,
            tf_model.trainable_weights
        )
    )

    tf_total_loss += float(
        loss_value.numpy()
    )

    tf_batches += 1

    # Show progress every 100 batches
    if (batch_index + 1) % 100 == 0:

        print(
            f"TensorFlow batch "
            f"{batch_index + 1}/{int(np.ceil(len(tf_x_train) / BATCH_SIZE))} "
            f"| Loss: {loss_value.numpy():.4f}"
        )

end_tf = time.perf_counter()

emissions_tf = tracker_tf.stop()

tf_duration = end_tf - start_tf

# Safety check in case CodeCarbon returns None
if emissions_tf is None:
    emissions_tf = 0.0

tf_average_loss = (
    tf_total_loss / tf_batches
)

print("\nTensorFlow completed.")

print(
    f"Execution time : "
    f"{tf_duration:.4f} seconds"
)

print(
    f"Average loss   : "
    f"{tf_average_loss:.4f}"
)

print(
    f"CO2 emissions  : "
    f"{emissions_tf:.8f} kg CO2eq"
)


# ==============================================================================
# 6. BENCHMARK COMPARISON
# ==============================================================================

print("\n" + "=" * 70)
print("BENCHMARK RESULTS")
print("=" * 70)

print(
    f"{'Metric':<25}"
    f"{'PyTorch':<20}"
    f"{'TensorFlow':<20}"
)

print("-" * 65)

print(
    f"{'Parameters':<25}"
    f"{pt_params:<20,}"
    f"{tf_params:<20,}"
)

print(
    f"{'Execution Time (s)':<25}"
    f"{pt_duration:<20.4f}"
    f"{tf_duration:<20.4f}"
)

print(
    f"{'Average Loss':<25}"
    f"{pt_average_loss:<20.4f}"
    f"{tf_average_loss:<20.4f}"
)

print(
    f"{'CO2 (kg)':<25}"
    f"{emissions_pt:<20.8f}"
    f"{emissions_tf:<20.8f}"
)


# ------------------------------------------------------------------------------
# Determine faster framework
# ------------------------------------------------------------------------------

if pt_duration < tf_duration:

    faster_framework = "PyTorch"

    speed_difference = (
        tf_duration - pt_duration
    )

else:

    faster_framework = "TensorFlow"

    speed_difference = (
        pt_duration - tf_duration
    )


# ------------------------------------------------------------------------------
# Determine framework with lower estimated emissions
# ------------------------------------------------------------------------------

if emissions_pt < emissions_tf:

    greener_framework = "PyTorch"

elif emissions_tf < emissions_pt:

    greener_framework = "TensorFlow"

else:

    greener_framework = "Equal"


print()

print(
    f"Faster framework: "
    f"{faster_framework}"
)

print(
    f"Execution-time difference: "
    f"{speed_difference:.4f} seconds"
)

print(
    f"Lower estimated emissions: "
    f"{greener_framework}"
)


# ==============================================================================
# 7. AUTOMATED WORD DOCUMENT GENERATION
# ==============================================================================

print("\n" + "=" * 70)
print("GENERATING WORD REPORT")
print("=" * 70)

doc = Document()


# ------------------------------------------------------------------------------
# Document Title
# ------------------------------------------------------------------------------

title = doc.add_heading(
    "Green AI Research: Framework Benchmarking Report",
    level=0
)

title.alignment = WD_ALIGN_PARAGRAPH.CENTER


# ------------------------------------------------------------------------------
# Metadata
# ------------------------------------------------------------------------------

doc.add_paragraph(
    f"Generated Automatically: "
    f"{time.strftime('%Y-%m-%d %H:%M:%S')}"
)

doc.add_paragraph(
    "Author: Research Team\n"
    "Target Frameworks: PyTorch vs. TensorFlow\n"
    "Dataset: CIFAR-10\n"
    "Training Duration: 1 Epoch\n"
    "Benchmark Device: CPU"
)


# ==============================================================================
# REPORT SECTION 1
# ==============================================================================

doc.add_heading(
    "1. Model Architecture & Parameter Parity",
    level=1
)

doc.add_paragraph(
    "To improve experimental comparability, both model architectures "
    "were matched layer-by-layer. Each implementation contains two "
    "convolutional blocks with 32 and 64 filters, Batch Normalization, "
    "ReLU activation, Max Pooling, Global Average Pooling, and Dense "
    "classification layers."
)


# ------------------------------------------------------------------------------
# Table 1: Model Parity
# ------------------------------------------------------------------------------

table_parity = doc.add_table(
    rows=3,
    cols=3
)

table_parity.style = "Table Grid"

hdr_cells = table_parity.rows[0].cells

hdr_cells[0].text = "Framework"
hdr_cells[1].text = "Trainable Parameters"
hdr_cells[2].text = "Parity Status"

row1 = table_parity.rows[1].cells

row1[0].text = "PyTorch"
row1[1].text = f"{pt_params:,}"
row1[2].text = parity_status

row2 = table_parity.rows[2].cells

row2[0].text = "TensorFlow / Keras"
row2[1].text = f"{tf_params:,}"
row2[2].text = parity_status


# ==============================================================================
# REPORT SECTION 2
# ==============================================================================

doc.add_heading(
    "2. Empirical Benchmark Results (1 Epoch)",
    level=1
)


# ------------------------------------------------------------------------------
# Table 2: Benchmark Results
# ------------------------------------------------------------------------------

table_results = doc.add_table(
    rows=3,
    cols=4
)

table_results.style = "Table Grid"

hdr_cells2 = table_results.rows[0].cells

hdr_cells2[0].text = "Framework"
hdr_cells2[1].text = "Execution Time (Seconds)"
hdr_cells2[2].text = "Average Training Loss"
hdr_cells2[3].text = "Estimated Emissions (kg CO2eq)"

res1 = table_results.rows[1].cells

res1[0].text = "PyTorch"
res1[1].text = f"{pt_duration:.4f} s"
res1[2].text = f"{pt_average_loss:.4f}"
res1[3].text = f"{emissions_pt:.8f} kg"

res2 = table_results.rows[2].cells

res2[0].text = "TensorFlow"
res2[1].text = f"{tf_duration:.4f} s"
res2[2].text = f"{tf_average_loss:.4f}"
res2[3].text = f"{emissions_tf:.8f} kg"


# ==============================================================================
# REPORT SECTION 3
# ==============================================================================

doc.add_heading(
    "3. Key Research Observations",
    level=1
)

doc.add_paragraph(
    f"Parameter Counts: PyTorch contained {pt_params:,} trainable "
    f"parameters and TensorFlow contained {tf_params:,} trainable "
    f"parameters. Parity status: {parity_status}.",
    style="List Bullet"
)

doc.add_paragraph(
    f"Execution Time: PyTorch required {pt_duration:.2f} seconds "
    f"and TensorFlow required {tf_duration:.2f} seconds.",
    style="List Bullet"
)

doc.add_paragraph(
    f"Execution Performance: {faster_framework} was faster during "
    f"this experimental run by approximately "
    f"{speed_difference:.2f} seconds.",
    style="List Bullet"
)

doc.add_paragraph(
    f"Estimated Carbon Emissions: PyTorch produced an estimate of "
    f"{emissions_pt:.8f} kg CO2eq while TensorFlow produced an "
    f"estimate of {emissions_tf:.8f} kg CO2eq.",
    style="List Bullet"
)

doc.add_paragraph(
    f"Lower Estimated Emissions: {greener_framework}.",
    style="List Bullet"
)

doc.add_paragraph(
    "CodeCarbon was used to estimate energy-related carbon emissions "
    "associated with each benchmark run.",
    style="List Bullet"
)


# ==============================================================================
# REPORT SECTION 4
# ==============================================================================

doc.add_heading(
    "4. Experimental Configuration",
    level=1
)

doc.add_paragraph(
    f"Random Seed: {SEED}\n"
    f"Batch Size: {BATCH_SIZE}\n"
    f"Optimizer: Adam\n"
    f"Learning Rate: {LEARNING_RATE}\n"
    f"Dataset: CIFAR-10\n"
    f"Training Samples: {len(pt_trainset):,}\n"
    f"Epochs: 1\n"
    f"Execution Device: CPU\n"
    f"PyTorch Version: {torch.__version__}\n"
    f"TensorFlow Version: {tf.__version__}"
)


# ==============================================================================
# 8. SAVE WORD DOCUMENT
# ==============================================================================

output_path = "Green_AI_Benchmark_Report.docx"

doc.save(
    output_path
)

print("\n" + "=" * 70)
print("SUCCESS!")
print(f"Word Document Generated: {output_path}")
print("=" * 70)